"""Small local HTTP server for the investment journal (no external web framework)."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import sys
import threading
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from decimal import InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse, urlencode
from urllib.request import Request, urlopen

from .ledger import calculate, dec, money
from .ai_analysis import analyze
from .daily import daily_profit


ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("INVESTMENT_DB", ROOT.parent / "data" / "investments.db"))
STATIC = ROOT / "static"
DEFAULT_ASSETS = [
    ("余额宝", "CASH", "", "cash", 0.01),
    ("纳斯达克100 ETF", "513100", "1.513100", "security", 0),
    ("标普500 ETF", "513500", "1.513500", "security", 0),
    ("沪深300 ETF", "510300", "1.510300", "security", 0),
    ("广发存金 · Au99.99参考价", "Au9999", "118.AU9999", "security", 0),
]


def derive_secid(code: str) -> str:
    """Resolve a mainland security or Au99.99 code to a quote ID."""
    code = code.strip()
    if code.upper() in {"AU9999", "AU99.99"}:
        return "118.AU9999"
    if not re.fullmatch(r"\d{6}", code):
        raise ValueError("请输入沪深市场的6位证券代码，例如 159941")
    if code == "000300":  # 沪深300指数；普通 0/1 开头的证券在深市
        return "1.000300"
    market = "1" if code[0] in "5679" else "0"
    return f"{market}.{code}"


def connection() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH, timeout=10)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
        CREATE TABLE IF NOT EXISTS assets (
            id INTEGER PRIMARY KEY, name TEXT NOT NULL, code TEXT NOT NULL,
            secid TEXT NOT NULL DEFAULT '', kind TEXT NOT NULL,
            annual_rate REAL NOT NULL DEFAULT 0, UNIQUE(code)
        );
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY, time TEXT NOT NULL, type TEXT NOT NULL,
            asset_id INTEGER REFERENCES assets(id), quantity REAL NOT NULL DEFAULT 0,
            price REAL NOT NULL DEFAULT 0, fee REAL NOT NULL DEFAULT 0,
            amount REAL NOT NULL, note TEXT NOT NULL DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS quotes (
            asset_id INTEGER PRIMARY KEY REFERENCES assets(id),
            price REAL NOT NULL, time TEXT NOT NULL, source TEXT NOT NULL
        );
    """)
    if "prev_close" not in {column["name"] for column in db.execute("PRAGMA table_info(quotes)")}:
        db.execute("ALTER TABLE quotes ADD COLUMN prev_close REAL")
    db.execute("CREATE TABLE IF NOT EXISTS asset_targets (asset_id INTEGER PRIMARY KEY REFERENCES assets(id) ON DELETE CASCADE, target_pct REAL NOT NULL CHECK(target_pct>=0 AND target_pct<=100))")
    db.commit()
    if "is_active" not in {column["name"] for column in db.execute("PRAGMA table_info(assets)")}:
        db.execute("ALTER TABLE assets ADD COLUMN is_active INTEGER NOT NULL DEFAULT 1")
        db.commit()
    if not db.execute("SELECT 1 FROM assets LIMIT 1").fetchone():
        db.executemany("INSERT INTO assets(name,code,secid,kind,annual_rate) VALUES(?,?,?,?,?)", DEFAULT_ASSETS)
        db.commit()
    # Repair ETF codes saved with a manually entered market prefix by older versions.
    for asset in db.execute("SELECT id,code,secid FROM assets WHERE kind='security'").fetchall():
        code = asset["code"]
        if re.fullmatch(r"\d{6}", code) and code.startswith(("15", "16", "51", "52", "53", "56", "58")):
            expected = derive_secid(code)
            if asset["secid"] != expected:
                db.execute("UPDATE assets SET secid=? WHERE id=?", (expected, asset["id"]))
                db.execute("DELETE FROM quotes WHERE asset_id=?", (asset["id"],))
                db.commit()
    return db


@contextmanager
def database():
    db = connection()
    try:
        yield db
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def rows(db: sqlite3.Connection, table: str) -> list[dict]:
    order = "time,id" if table == "transactions" else ("asset_id" if table == "quotes" else "id")
    return [dict(r) for r in db.execute(f"SELECT * FROM {table} ORDER BY {order}")]


def current_state(db: sqlite3.Connection) -> dict:
    assets = rows(db, "assets")
    transactions = rows(db, "transactions")
    quotes = {q["asset_id"]: q for q in rows(db, "quotes")}
    now=datetime.now()
    portfolio = calculate(assets, transactions, quotes, now)
    targets={str(r["asset_id"]):r["target_pct"] for r in db.execute("SELECT asset_id,target_pct FROM asset_targets")}
    day=daily_profit(assets,transactions,quotes,now)
    return {"targets":targets,"daily":day,"initialized": bool(transactions), "assets": assets,
            "transactions": transactions, "quotes": list(quotes.values()),
            "portfolio": portfolio}


def parse_time(value: str) -> str:
    at = datetime.fromisoformat(value)
    if at > datetime.now():
        raise ValueError("交易时间不能晚于当前时间")
    return at.isoformat(timespec="minutes")


def positive(value: object, label: str, allow_zero: bool = False) -> float:
    try:
        number = dec(value)
    except (InvalidOperation, TypeError):
        raise ValueError(f"请填写有效的{label}") from None
    if not number.is_finite() or number < 0 or (number == 0 and not allow_zero):
        raise ValueError(f"{label}必须为正数")
    return float(number)


def validate_history(db: sqlite3.Connection) -> None:
    state = current_state(db)
    qty = {a["id"]: 0.0 for a in state["assets"]}
    for tx in state["transactions"]:
        if tx["type"] == "buy":
            qty[tx["asset_id"]] += tx["quantity"]
        elif tx["type"] == "sell":
            qty[tx["asset_id"]] -= tx["quantity"]
            if qty[tx["asset_id"]] < -1e-7:
                raise ValueError("卖出数量超过当时持仓")
    for point in state["portfolio"]["timeline"]:
        if point["cash"] < -0.01:
            raise ValueError("交易使现金余额为负，请检查时间或金额")


def transaction_values(db: sqlite3.Connection, body: dict) -> tuple:
    kind = body.get("type")
    if kind not in {"deposit", "withdraw", "buy", "sell"}:
        raise ValueError("不支持的交易类型")
    at = parse_time(str(body.get("time", "")))
    note = str(body.get("note", ""))[:300]
    if kind in {"deposit", "withdraw"}:
        asset_id, quantity, price, fee = None, 0, 0, 0
        amount = positive(body.get("amount"), "金额")
    else:
        asset_id = int(body.get("asset_id"))
        asset = db.execute("SELECT * FROM assets WHERE id=?", (asset_id,)).fetchone()
        if not asset or asset["kind"] != "security":
            raise ValueError("请选择有效的证券标的")
        price = positive(body.get("price"), "价格")
        fee = positive(body.get("fee", 0), "手续费", allow_zero=True)
        if body.get("input_mode") == "amount" and kind == "buy":
            budget = dec(positive(body.get("amount"), "买入金额"))
            if budget <= dec(fee):
                raise ValueError("买入金额必须大于手续费")
            quantity = float(((budget - dec(fee)) / dec(price)).quantize(dec("0.00000001")))
        else:
            quantity = positive(body.get("quantity"), "数量")
        if quantity <= 0:
            raise ValueError("金额过小，无法生成有效数量")
        gross = money(dec(price) * dec(quantity))
        amount = float(gross + dec(fee) if kind == "buy" else gross - dec(fee))
        if amount <= 0:
            raise ValueError("成交金额必须大于手续费")
    return at, kind, asset_id, quantity, price, fee, amount, note


def add_transaction(db: sqlite3.Connection, body: dict) -> int:
    values = transaction_values(db, body)
    cursor = db.execute(
        "INSERT INTO transactions(time,type,asset_id,quantity,price,fee,amount,note) VALUES(?,?,?,?,?,?,?,?)",
        values,
    )
    try:
        validate_history(db)
    except Exception:
        db.rollback()
        raise
    db.commit()
    return cursor.lastrowid


def update_transaction(db: sqlite3.Connection, body: dict) -> dict:
    tx_id = int(body.get("id"))
    if not db.execute("SELECT 1 FROM transactions WHERE id=?", (tx_id,)).fetchone():
        raise ValueError("交易记录不存在")
    values = transaction_values(db, body)
    db.execute(
        "UPDATE transactions SET time=?,type=?,asset_id=?,quantity=?,price=?,fee=?,amount=?,note=? WHERE id=?",
        (*values, tx_id),
    )
    try:
        validate_history(db)
    except Exception:
        db.rollback()
        raise
    db.commit()
    return current_state(db)


def fetch_quote_eastmoney(asset: dict) -> dict:
    secid = asset["secid"]
    if not re.fullmatch(r"[0-9]{1,3}\.[A-Za-z0-9]{1,20}", secid):
        raise ValueError("行情代码格式错误")
    params = urlencode({"secid": secid, "fields": "f43,f57,f58,f59,f60,f86"})
    req = Request("https://push2.eastmoney.com/api/qt/stock/get?" + params,
                  headers={"User-Agent": "Mozilla/5.0", "Referer": "https://quote.eastmoney.com/"})
    with urlopen(req, timeout=5) as response:
        payload = json.load(response)
    data = payload.get("data") or {}
    raw = data.get("f43")
    precision = data.get("f59")
    if raw in (None, "-"):
        raise ValueError("无最新价格")
    price = float(raw)
    if isinstance(raw, int) and isinstance(precision, int) and 0 <= precision <= 6:
        price /= 10 ** precision
    if price <= 0:
        raise ValueError("无有效价格")
    stamp = str(data.get("f86") or "")
    quote_time = None
    if re.fullmatch(r"\d{10}", stamp):
        quote_time = datetime.fromtimestamp(int(stamp), timezone(timedelta(hours=8))).replace(tzinfo=None).isoformat(timespec="minutes")
    elif re.fullmatch(r"\d{12}", stamp):
        quote_time = datetime.strptime(stamp, "%Y%m%d%H%M").isoformat(timespec="minutes")
    if quote_time is None: raise ValueError("行情缺少报价时间")
    previous=data.get("f60")
    try:
        previous=float(previous)
        if isinstance(data.get("f60"), int) and isinstance(precision,int) and 0<=precision<=6: previous/=10**precision
        if previous<=0: previous=None
    except (ValueError,TypeError): previous=None
    return {"asset_id": asset["id"], "price": price,
            "prev_close":previous, "time": quote_time, "source": "东方财富"}


def fetch_quote_sge(asset: dict) -> dict:
    """Read the Shanghai Gold Exchange's published delayed Au99.99 quote."""
    req = Request("https://www.sge.com.cn/graph/quotations",
                  data=urlencode({"instid": "Au99.99"}).encode("ascii"),
                  headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.sge.com.cn/sjzx/yshqbg",
                           "Content-Type": "application/x-www-form-urlencoded",
                           "X-Requested-With": "XMLHttpRequest"})
    with urlopen(req, timeout=8) as response:
        payload = json.load(response)
    if payload.get("heyue") != "Au99.99":
        raise ValueError("黄金合约代码不匹配")
    match = re.fullmatch(r"(\d{4})年(\d{1,2})月(\d{1,2})日 (\d{2}):(\d{2}):(\d{2})", str(payload.get("delaystr", "")))
    if not match:
        raise ValueError("黄金行情缺少报价时间")
    reported = datetime(*map(int, match.groups()))
    times, prices = payload.get("times") or [], payload.get("data") or []
    for minute, raw in reversed(list(zip(times, prices))):
        try:
            price = float(raw)
            hour, minute_number = map(int, minute.split(":"))
            quote_time = reported.replace(hour=hour, minute=minute_number, second=0)
            if quote_time > reported:
                quote_time -= timedelta(days=1)
            if (price > 0 and 0 <= (reported - quote_time).total_seconds() <= 24 * 3600
                    and quote_time <= datetime.now() + timedelta(minutes=5)):
                return {"asset_id": asset["id"], "price": price,
                        "time": quote_time.isoformat(timespec="minutes"), "source": "上金所延时行情"}
        except (ValueError, TypeError):
            continue
    raise ValueError("黄金行情暂无有效价格")


def fetch_quote_tencent(asset: dict) -> dict:
    secid = asset["secid"]
    if not re.fullmatch(r"[01]\.\d{6}", secid):
        raise ValueError("该行情代码没有备用报价源")
    market, code = secid.split(".")
    symbol = ("sh" if market == "1" else "sz") + code
    req = Request(f"https://qt.gtimg.cn/q={symbol}",
                  headers={"User-Agent": "Mozilla/5.0", "Referer": "https://gu.qq.com/"})
    with urlopen(req, timeout=5) as response:
        payload = response.read().decode("gbk", errors="replace")
    fields = payload.partition('"')[2].split("~")
    if len(fields) < 31 or fields[2] != code:
        raise ValueError("备用行情无报价")
    price = positive(fields[3], "报价")
    stamp = fields[30]
    if not re.fullmatch(r"\d{14}", stamp):
        raise ValueError("备用行情无报价时间")
    quote_time = datetime.strptime(stamp, "%Y%m%d%H%M%S").isoformat(timespec="minutes")
    try: previous=positive(fields[4], "昨收")
    except ValueError: previous=None
    return {"asset_id": asset["id"], "price": price,
            "prev_close":previous, "time": quote_time, "source": "腾讯行情"}


def fetch_quote(asset: dict) -> dict:
    if asset["code"].upper() in {"AU9999", "AU99.99"}:
        try:
            return fetch_quote_sge(asset)
        except Exception:
            return fetch_quote_eastmoney(asset)
    try:
        return fetch_quote_eastmoney(asset)
    except Exception:
        return fetch_quote_tencent(asset)


def save_quote(db: sqlite3.Connection, quote: dict) -> None:
    db.execute("INSERT INTO quotes(asset_id,price,time,source,prev_close) VALUES(?,?,?,?,?) "
               "ON CONFLICT(asset_id) DO UPDATE SET price=excluded.price,time=excluded.time,source=excluded.source,prev_close=COALESCE(excluded.prev_close,CASE WHEN substr(quotes.time,1,10)=substr(excluded.time,1,10) THEN quotes.prev_close ELSE NULL END)",
               (quote["asset_id"], quote["price"], quote["time"], quote["source"],quote.get("prev_close")))


def add_asset(db: sqlite3.Connection, body: dict) -> dict:
    name = str(body.get("name", "")).strip()[:60]
    code = str(body.get("code", "")).strip()
    if code.upper() in {"AU9999", "AU99.99"}:
        code = "Au9999"
    if not name:
        raise ValueError("请填写标的名称")
    secid = derive_secid(code)
    existing = db.execute("SELECT id,is_active FROM assets WHERE code=?", (code,)).fetchone()
    if existing and existing["is_active"]:
        raise ValueError("该标的代码已添加")
    if existing:
        db.execute("UPDATE assets SET is_active=1,name=?,secid=? WHERE id=?", (name, secid, existing["id"]))
        asset_id = existing["id"]
    else:
        cursor = db.execute("INSERT INTO assets(name,code,secid,kind) VALUES(?,?,?,'security')", (name, code, secid))
        asset_id = cursor.lastrowid
    db.commit()
    asset = {"id": asset_id, "name": name, "code": code, "secid": secid}
    warning = None
    try:
        save_quote(db, fetch_quote(asset))
        db.commit()
    except Exception as exc:
        warning = f"标的已添加，但暂时无法获取行情：{exc}"
    result = current_state(db)
    if warning:
        result["quote_warning"] = warning
    return result


def delete_asset(db: sqlite3.Connection, asset_id: int) -> dict:
    asset = db.execute("SELECT * FROM assets WHERE id=?", (asset_id,)).fetchone()
    if not asset or asset["kind"] != "security" or not asset["is_active"]:
        raise ValueError("标的不存在或已移除")
    used = db.execute("SELECT 1 FROM transactions WHERE asset_id=? LIMIT 1", (asset_id,)).fetchone()
    if used:
        # Preserve the investment ledger; held assets remain visible until sold.
        db.execute("UPDATE assets SET is_active=0 WHERE id=?", (asset_id,))
    else:
        db.execute("DELETE FROM quotes WHERE asset_id=?", (asset_id,))
        db.execute("DELETE FROM assets WHERE id=?", (asset_id,))
    db.commit()
    return current_state(db)


def restore_asset(db: sqlite3.Connection, asset_id: int) -> dict:
    asset = db.execute("SELECT kind,is_active FROM assets WHERE id=?", (asset_id,)).fetchone()
    if not asset or asset["kind"] != "security" or asset["is_active"]:
        raise ValueError("标的不存在或已在关注列表中")
    db.execute("UPDATE assets SET is_active=1 WHERE id=?", (asset_id,))
    db.commit()
    return current_state(db)


def refresh_quotes(db: sqlite3.Connection) -> dict:
    state = current_state(db)
    held = {p["id"] for p in state["portfolio"]["positions"] if p["quantity"] > 0}
    assets = [a for a in state["assets"] if a["kind"] == "security" and a["secid"] and
              (a["is_active"] or a["id"] in held)]
    errors = []
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = {pool.submit(fetch_quote, a): a for a in assets}
        for future, asset in futures.items():
            try:
                q = future.result()
                save_quote(db, q)
            except Exception as exc:
                errors.append(f"{asset['name']}: {exc}")
    db.commit()
    return {"state": current_state(db), "errors": errors}


def export_file(db: sqlite3.Connection, fmt: str) -> tuple[bytes, str, str]:
    state = current_state(db)
    assets = {a["id"]: a for a in state["assets"]}
    headers = ["时间", "类型", "标的", "代码", "价格", "数量", "手续费", "现金流金额",
               "交易后标的持仓", "交易后标的收益", "交易后标的收益率%", "交易后账户总资产",
               "交易后账户收益", "交易后账户收益率%", "备注"]
    kinds = {"deposit": "入金", "withdraw": "出金", "buy": "买入", "sell": "卖出"}
    snapshots = {p["tx_id"]: p for p in state["portfolio"]["timeline"] if p["tx_id"] is not None}
    tx_rows = []
    for t in state["transactions"]:
        point = snapshots[t["id"]]
        position = point["asset_states"].get(str(t["asset_id"])) if t["asset_id"] else None
        tx_rows.append([t["time"], kinds[t["type"]], assets[t["asset_id"]]["name"] if t["asset_id"] else "余额宝",
                        assets[t["asset_id"]]["code"] if t["asset_id"] else "CASH", t["price"], t["quantity"],
                        t["fee"], t["amount"], position["quantity"] if position else None,
                        position["profit"] if position else None, position["return_rate"] if position else None,
                        point["equity"], point["profit"], point["return_rate"], t["note"]])
    if fmt == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(headers)
        writer.writerows(tx_rows)
        return ("\ufeff" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "investment_trades.csv"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "交易流水"
        ws.append(headers)
        for row in tx_rows:
            ws.append(row)
        summary = wb.create_sheet("当前持仓")
        summary.append(["标的", "代码", "数量", "成本", "最新价", "市值", "已实现收益", "未实现收益", "总收益", "收益率%"])
        summary.append(["余额宝", "CASH", None, state["portfolio"]["cash"] - state["portfolio"]["interest"],
                        None, state["portfolio"]["cash"], state["portfolio"]["interest"],
                        None, state["portfolio"]["interest"], None])
        for p in state["portfolio"]["positions"]:
            summary.append([p[k] for k in ("name", "code", "quantity", "cost", "price", "market_value", "realized", "unrealized", "profit", "return_rate")])
        curve = wb.create_sheet("资产曲线")
        curve.append(["时间", "总资产", "累计收益", "现金", "累计净投入"])
        for p in state["portfolio"]["timeline"]:
            curve.append([p[k] for k in ("time", "equity", "profit", "cash", "contributions")])
        asset_curve = wb.create_sheet("标的曲线")
        asset_curve.append(["时间", "标的", "代码", "持仓数量", "持仓市值", "已实现收益", "浮动盈亏", "累计盈亏", "收益率%"])
        for point in state["portfolio"]["timeline"]:
            for asset in state["assets"]:
                if asset["kind"] != "security":
                    continue
                value = point["asset_states"].get(str(asset["id"]), {})
                asset_curve.append([point["time"], asset["name"], asset["code"], value.get("quantity", 0),
                                    value.get("value", 0), value.get("realized", 0), value.get("unrealized", 0),
                                    value.get("profit", 0), value.get("return_rate", 0)])
        for sheet in wb:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="1D3444")
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 3, 12), 34)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "investment_report.xlsx"
    raise ValueError("未知导出格式")


class Handler(BaseHTTPRequestHandler):
    def send_json(self, value: object, status: int = 200) -> None:
        data = json.dumps(value, ensure_ascii=False, allow_nan=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self) -> None:
        path = urlparse(self.path).path
        if path == "/api/state":
            with database() as db:
                self.send_json(current_state(db))
            return
        if path.startswith("/api/export/"):
            try:
                with database() as db:
                    data, mime, name = export_file(db, path.rsplit("/", 1)[-1])
                self.send_response(200)
                self.send_header("Content-Type", mime)
                self.send_header("Content-Disposition", f'attachment; filename="{name}"')
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
            except (ValueError, ImportError) as exc:
                self.send_json({"error": str(exc)}, 400)
            return
        static_files = {"/": ("index.html", "text/html"), "/app.js": ("app.js", "text/javascript"),
                        "/style.css": ("style.css", "text/css")}
        if path not in static_files:
            self.send_error(404)
            return
        filename, mime = static_files[path]
        data = (STATIC / filename).read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", mime + "; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_POST(self) -> None:
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length > 100_000:
                raise ValueError("请求内容过大")
            body = json.loads(self.rfile.read(length) or b"{}")
            path = urlparse(self.path).path
            with database() as db:
                if path == "/api/ai/analyze":
                    origin = self.headers.get("Origin")
                    if origin != "http://" + self.headers.get("Host", ""):
                        raise ValueError("无效的请求来源")
                    result = analyze(current_state(db), body)
                elif path == "/api/initialize":
                    if rows(db, "transactions"):
                        raise ValueError("初始资金已设置")
                    body = {"type": "deposit", "time": body.get("time"), "amount": body.get("amount"), "note": "初始资金"}
                    add_transaction(db, body)
                    result = current_state(db)
                elif path == "/api/transactions":
                    add_transaction(db, body)
                    result = current_state(db)
                elif path == "/api/transactions/update":
                    result = update_transaction(db, body)
                elif path == "/api/targets":
                    db.execute("BEGIN IMMEDIATE")
                    asset_id=int(body.get("asset_id"))
                    if not db.execute("SELECT 1 FROM assets WHERE id=? AND kind='security'",(asset_id,)).fetchone(): raise ValueError("标的不存在")
                    value=body.get("target_pct")
                    if value in ("",None):
                        db.execute("DELETE FROM asset_targets WHERE asset_id=?",(asset_id,))
                    else:
                        target=positive(value,"目标比例",allow_zero=True)
                        if target>100: raise ValueError("目标比例不能超过100%")
                        other=db.execute("SELECT COALESCE(SUM(target_pct),0) FROM asset_targets WHERE asset_id<>?",(asset_id,)).fetchone()[0]
                        if other+target>100.0000001: raise ValueError("所有标的目标比例合计不能超过100%")
                        db.execute("INSERT INTO asset_targets(asset_id,target_pct) VALUES(?,?) ON CONFLICT(asset_id) DO UPDATE SET target_pct=excluded.target_pct",(asset_id,target))
                    db.commit()
                    result=current_state(db)
                elif path == "/api/assets":
                    result = add_asset(db, body)
                elif path == "/api/assets/delete":
                    result = delete_asset(db, int(body.get("id")))
                elif path == "/api/assets/restore":
                    result = restore_asset(db, int(body.get("id")))
                elif path == "/api/quotes/refresh":
                    result = refresh_quotes(db)
                elif path == "/api/quotes/manual":
                    asset_id = int(body.get("asset_id"))
                    asset = db.execute("SELECT kind FROM assets WHERE id=?", (asset_id,)).fetchone()
                    if not asset or asset["kind"] != "security":
                        raise ValueError("无效标的")
                    price = positive(body.get("price"), "价格")
                    stamp = parse_time(str(body.get("time", "")))
                    prev=body.get("prev_close")
                    prev=positive(prev,"昨收价") if prev not in ("",None) else None
                    save_quote(db,{"asset_id":asset_id,"price":price,"time":stamp,"source":"手动","prev_close":prev})
                    db.commit()
                    result = current_state(db)
                elif path == "/api/settings/cash-rate":
                    rate = positive(body.get("annual_rate"), "年化利率", allow_zero=True)
                    if rate > 1:
                        raise ValueError("年化利率不能超过100%")
                    db.execute("UPDATE assets SET annual_rate=? WHERE kind='cash'", (rate,))
                    db.commit()
                    result = current_state(db)
                elif path == "/api/transactions/delete":
                    tx_id = int(body.get("id"))
                    db.execute("DELETE FROM transactions WHERE id=?", (tx_id,))
                    validate_history(db)
                    db.commit()
                    result = current_state(db)
                else:
                    self.send_error(404)
                    return
            self.send_json(result)
        except (ValueError, TypeError, KeyError, sqlite3.IntegrityError) as exc:
            self.send_json({"error": str(exc)}, 400)
        except Exception as exc:
            self.send_json({"error": f"服务器错误: {exc}"}, 500)


def serve(port: int = 8766, open_browser: bool = False) -> None:
    with database():
        pass
    server = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    url = f"http://127.0.0.1:{port}"
    print(f"投资记录工具已启动：{url}", flush=True)
    if open_browser:
        threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    server.serve_forever()


if __name__ == "__main__":
    serve(int(os.environ.get("INVESTMENT_PORT", "8766")), "--open" in sys.argv)
